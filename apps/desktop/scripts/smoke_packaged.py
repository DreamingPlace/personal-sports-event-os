"""Exercise the bundled ticket book sidecar with an empty PATH; Python here is only the test runner.

Usage: smoke_packaged.py <sidecar binary> [--require-ocr]
All data is invented (the built-in example book).
"""
import json
import os
import select
import subprocess
import sys
import tempfile
from pathlib import Path

binary = Path(sys.argv[1]).resolve()
require_ocr = '--require-ocr' in sys.argv
checks = []


def start():
    return subprocess.Popen([str(binary)], stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True,
                            env={'PATH': '/nonexistent', 'HOME': os.environ['HOME'], 'LANG': 'en_US.UTF-8', 'TMPDIR': tempfile.gettempdir()})


with tempfile.TemporaryDirectory(prefix='ticketbook-packaged-') as workspace:
    process = start()
    counter = 0

    def call(method, timeout=120, **params):
        global counter
        counter += 1
        rid = str(counter)
        process.stdin.write(json.dumps({'id': rid, 'method': method, 'params': params}) + '\n')
        process.stdin.flush()
        if not select.select([process.stdout], [], [], timeout)[0]:
            raise TimeoutError(method)
        response = json.loads(process.stdout.readline())
        assert response['id'] == rid, response
        return response

    def ok(method, **params):
        response = call(method, **params)
        assert response['ok'], response
        return response['result']

    try:
        health = ok('health')
        checks.append(f"health {health['version']} / desktop {health['desktop_version']} / ocr {health['ocr']}")
        if require_ocr:
            assert health['ocr'], 'OCR was not bundled'
        path = str(Path(workspace) / 'example.ticketbook')
        state = ok('create_book', path=path, password='pw', example=True)
        assert state['ledger']['ok'], state['ledger']['problems']
        checks.append('create encrypted example book')
        state = ok('edit', ops=[{'op': 'set', 'path': ['prices', 'pre', 'VIP'], 'value': '550'}])
        assert state['book']['prices']['pre']['VIP'] == '550'
        checks.append('edit')
        out = ok('export_inventory', path=str(Path(workspace) / 'count.xlsx'))
        assert Path(out['path']).stat().st_size > 0
        checks.append('export inventory count (Excel)')
        text = '\t100000001\t示例杯\t数量（张）\t2000\t500\t20\t1500\t25%\n\t金额（元）\t200000\t50000\t2000\t150000\n'
        state = ok('damai_add', text=text, at='2030-11-20T10:00')
        assert state['damai']['latest']['total']['sold_qty'] == 500
        checks.append('read pasted Damai table')
        import docx  # test runner only, to write a template
        template = docx.Document()
        template.add_paragraph('共{{场次数}}场')
        template.add_paragraph('{{票价表}}')
        template.save(Path(workspace) / 'template.docx')
        state = ok('template_add', path=str(Path(workspace) / 'template.docx'))
        made = ok('report_make', id=state['book']['templates'][0]['id'], path=str(Path(workspace) / 'report.docx'))
        assert '共5场' in '\n'.join(p.text for p in docx.Document(made['path']).paragraphs), made
        checks.append('fill a Word report template')
        from openpyxl import Workbook  # test runner only, to write an invented seat sheet
        wb = Workbook()
        wb.active.cell(row=1, column=2, value='东A区')
        for r in range(2, 5):
            for c in range(2, 8):
                wb.active.cell(row=r, column=c, value=c - 1)
        wb.save(Path(workspace) / 'seats.xlsx')
        zones = ok('seatmap_read', sheet=str(Path(workspace) / 'seats.xlsx'))['zones']
        assert zones == [{'name': '东A区', 'seats': 18, 'colour': None, 'colours': {}, 'sure': False}], zones
        from PIL import Image  # a PDF without text: the PDF reader must load and say so
        Image.new('RGB', (200, 200), (255, 255, 255)).save(Path(workspace) / 'blank.pdf')
        response = call('seatmap_read', sheet=str(Path(workspace) / 'blank.pdf'))
        assert not response['ok'] and '没有文字' in response['error']['message'], response
        checks.append('read seat sheets (Excel, PDF)')
        wb = Workbook()
        wb.active.append(['票品', '预赛', '决赛'])
        wb.active.append(['VIP', 555, 999])
        wb.active.append(['A档', 444, 888])
        wb.save(Path(workspace) / 'plan.xlsx')
        found = ok('setup_scan', paths=[str(Path(workspace) / 'plan.xlsx')])
        assert found['prices'][0]['options'][0]['prices']['决赛']['VIP'] == '999', found
        checks.append('recognise a planning sheet')
        if health['ocr']:
            from PIL import Image, ImageDraw, ImageFont  # test runner only, to draw a picture to read
            picture = Image.new('RGB', (900, 200), (10, 20, 70))
            ImageDraw.Draw(picture).text((40, 60), 'S3  2030-12-02 15:30:05', fill=(235, 235, 255), font=ImageFont.load_default(size=48))
            picture.save(Path(workspace) / 'screen.png')
            result = ok('live_read', paths=[str(Path(workspace) / 'screen.png')])['results'][0]
            assert result.get('session') == 'S3' and result['entry'].get('at') == '2030-12-02T15:30', result
            checks.append('screenshot reader (OCR model) works')
        process.stdin.close()
        process.wait(timeout=10)
        process = start()
        assert call('open_book', path=path, password='wrong')['error']['code'] == 'PASSWORD'
        state = ok('open_book', path=path, password='pw')
        assert state['book']['prices']['pre']['VIP'] == '550'
        checks.append('restart and reopen with password')
    finally:
        if process.poll() is None:
            process.stdin.close()
            process.wait(timeout=10)
print(json.dumps({'status': 'PASS', 'binary': str(binary), 'checks': checks}, ensure_ascii=False, indent=2))
