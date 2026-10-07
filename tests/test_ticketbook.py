"""Ticket book engine. All figures are invented."""
import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from ticketbook.desktop import Session
from ticketbook.edits import apply, compare, copy_event
from ticketbook.example import example_book
from ticketbook.export import forecast_table, inventory_count
from ticketbook.forecast import candidate, forecast, preview
from ticketbook.ledger import as_of, compute
from ticketbook.ocr import Box, parse_boxes
from ticketbook.live import damai_summary, live_summary, parse_damai_table
from ticketbook.model import BookError, new_book, normalize, price
from ticketbook.store import BookFile, WrongPassword


def tiny_book():
    """One session, one tier, round numbers, so every line can be checked by hand."""
    b = new_book('T', 2030)
    b['tiers'] = [{'code': 'A'}]
    b['bands'] = [{'code': 'x'}]
    b['prices'] = {'x': {'A': '100'}}
    b['layouts'] = [{'code': 'L', 'seats': {'A': 1000}}]
    b['sessions'] = [{'code': 'S1', 'date': '2030-01-01', 'band': 'x', 'layout': 'L'}]
    b['buckets'] = [
        {'id': 'h', 'name': '转播', 'kind': 'hold', 'default': {'A': 50}},
        {'id': 'c', 'name': '权益', 'kind': 'comp', 'default': {'A': 100}},
        {'id': 'p', 'name': '优先购', 'kind': 'priority', 'default': {'A': 150}, 'cap': {'A': 120}},
        {'id': 'r', 'name': '预留', 'kind': 'reserve', 'default': {'A': 200}},
    ]
    b['rounds'] = [{'code': 'R1', 'share': {'x': '40'}}, {'code': 'R2', 'share': {'x': '60'}}]
    b['sales'] = {'R1': {'S1': {'A': 180}}, 'R2': {'S1': {'A': 120}}}
    return normalize(b)


class LedgerTest(unittest.TestCase):
    def test_lines_add_up(self):
        r = compute(tiny_book())['sessions'][0]['tiers']['A']
        self.assertEqual((r['seats'], r['hold'], r['sellable']), (1000, 50, 950))
        self.assertEqual((r['comp'], r['priority'], r['reserve'], r['public']), (100, 150, 200, 500))
        self.assertEqual((r['sold'], r['left']), (300, 200))
        self.assertEqual(r['by_round'], {'R1': 180, 'R2': 120})

    def test_release_plan_follows_round_shares(self):
        s = compute(tiny_book())['sessions'][0]
        self.assertEqual(s['release_plan']['A'], {'R1': 200, 'R2': 300})

    def test_cap_is_a_warning(self):
        problems = compute(tiny_book())['problems']
        self.assertEqual([p['level'] for p in problems], ['warning'])
        self.assertIn('超过上限 120', problems[0]['message'])

    def test_over_allocation_is_an_error_naming_session_and_tier(self):
        b = apply(tiny_book(), [{'op': 'set', 'path': ['buckets', 'r', 'default', 'A'], 'value': 800}])
        result = compute(b)
        self.assertFalse(result['ok'])
        self.assertTrue(any('S1 A 分配超出可售座席 100 张' in p['message'] for p in result['problems']))

    def test_override_beats_default_and_band_filter(self):
        b = example_book()
        b['buckets'][2]['overrides'] = {'S2': {'VIP': 7}}
        b['buckets'][3]['bands'] = ['final']
        rows = {s['code']: s['tiers'] for s in compute(b)['sessions']}
        self.assertEqual(rows['S1']['VIP']['by_bucket']['partner'], 100)
        self.assertEqual(rows['S2']['VIP']['by_bucket']['partner'], 7)
        self.assertNotIn('venue', rows['S1']['A']['by_bucket'])
        self.assertEqual(rows['S5']['A']['by_bucket']['venue'], 30)

    def test_blocked_tier_price_is_base_minus_discount(self):
        b = example_book()
        self.assertEqual(price(b, 'rr', 'B遮挡'), Decimal('300'))
        b['prices']['rr']['B遮挡'] = '250'
        self.assertEqual(price(b, 'rr', 'B遮挡'), Decimal('250'))

    def test_missing_price_and_unknown_codes_are_reported(self):
        b = example_book()
        del b['prices']['final']['C']
        b['sessions'][0]['band'] = 'nope'
        messages = [p['message'] for p in compute(b)['problems']]
        self.assertTrue(any('S5 C 没有票价' in m for m in messages))
        self.assertTrue(any('S1 的价格段 nope 不存在' in m for m in messages))

    def test_day_pass_uses_seats_in_every_session_of_its_days(self):
        rows = {s['code']: s['tiers']['B'] for s in compute(example_book())['sessions']}
        self.assertTrue(all(r['product'] == 50 for r in rows.values()))

    def test_as_of_counts_only_given_rounds(self):
        r = compute(as_of(tiny_book(), ['R1']))['sessions'][0]['tiers']['A']
        self.assertEqual(r['sold'], 180)
        with self.assertRaises(BookError):
            as_of(tiny_book(), ['R9'])


class ForecastTest(unittest.TestCase):
    def test_paid_seats_times_price_times_fill(self):
        b = tiny_book()
        b['forecast']['band_fill'] = {'x': '80'}
        b['event']['agent_fee_pct'] = '5'
        f = forecast(b)
        # paid seats = public 500 + priority 150 + reserve 200 = 850 → 85,000 full, 68,000 at 80%
        self.assertEqual(Decimal(f['sessions'][0]['full']), Decimal(85000))
        self.assertEqual(Decimal(f['gross']), Decimal(68000))
        self.assertEqual(Decimal(f['agent_fee']), Decimal(3400))
        self.assertEqual(Decimal(f['net']), Decimal(64600))

    def test_china_method_uses_session_flag(self):
        b = example_book()
        f = {s['code']: s['fill'] for s in forecast(b)['sessions']}
        self.assertEqual((f['S1'], f['S2']), ('65', '95'))
        b['forecast']['session_fill'] = {'S1': '50'}
        self.assertEqual(forecast(b)['sessions'][0]['fill'], '50')

    def test_day_pass_revenue_is_per_day_sum(self):
        p = forecast(example_book())['products'][0]
        # day 1: 300+300, day 2: 400+400, day 3: 800 → (600+800+800) × 50
        self.assertEqual(Decimal(p['full']), Decimal(110000))

    def test_preview_never_changes_the_book(self):
        b = example_book()
        before = json.dumps(b, sort_keys=True)
        result = preview(b, [{'path': ['prices', 'pre', 'VIP'], 'value': '600'}, {'path': ['forecast', 'other_fill'], 'value': '70'}])
        self.assertEqual(json.dumps(b, sort_keys=True), before)
        self.assertEqual([c['path'] for c in result['changes']], [['prices', 'pre', 'VIP'], ['forecast', 'other_fill']])
        self.assertGreater(Decimal(result['after']['gross']), Decimal(result['before']['gross']))

    def test_only_forecast_inputs_can_change(self):
        with self.assertRaises(BookError):
            candidate(example_book(), [{'path': ['sessions', 'S1', 'date'], 'value': '2031-01-01'}])
        with self.assertRaises(BookError):
            candidate(example_book(), [{'path': ['prices', 'pre', 'VIP'], 'value': '-1'}])


class EditTest(unittest.TestCase):
    def test_rename_tier_moves_every_reference(self):
        b = apply(example_book(), [{'op': 'rename', 'list': 'tiers', 'key': 'B', 'to': 'B1'}])
        self.assertIn('B1', b['prices']['pre'])
        self.assertEqual(b['tiers'][3]['blocked_of'], 'B1')
        self.assertEqual(b['products'][0]['tier'], 'B1')
        self.assertTrue(compute(b)['ok'])

    def test_remove_session_drops_its_overrides_and_sales(self):
        b = tiny_book()
        b['buckets'][0]['overrides'] = {'S1': {'A': 1}}
        b = apply(b, [{'op': 'remove', 'list': 'sessions', 'key': 'S1'}])
        self.assertEqual(b['buckets'][0]['overrides'], {})
        self.assertEqual(b['sales']['R1'], {})

    def test_bad_numbers_are_refused(self):
        for value in ('abc', -5, '1.5'):
            with self.assertRaises(BookError):
                apply(tiny_book(), [{'op': 'set', 'path': ['layouts', 'L', 'seats', 'A'], 'value': value}])

    def test_compare_lists_changed_values(self):
        a = tiny_book()
        b = apply(a, [{'op': 'set', 'path': ['prices', 'x', 'A'], 'value': '120'}])
        self.assertEqual(compare(a, b), [{'path': ['prices', 'x', 'A'], 'old': '100', 'new': '120'}])

    def test_copy_event_moves_dates_and_clears_sales(self):
        b = copy_event(example_book(), '下一届', 2031, 365)
        self.assertEqual(b['sessions'][0]['date'], '2031-12-01')
        self.assertEqual(b['rounds'][0]['opens'], '2031-11-10T15:00')
        self.assertEqual(b['sales'], {})


class StoreTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.path = Path(self.dir.name) / 'event'

    def tearDown(self):
        self.dir.cleanup()

    def test_file_is_encrypted_and_reopens_only_with_password(self):
        f = BookFile.create(self.path, 'secret', example_book())
        raw = f.path.read_bytes()
        self.assertTrue(f.path.name.endswith('.ticketbook'))
        self.assertNotIn('示例'.encode('utf-8'), raw)
        self.assertNotIn(b'VIP', raw)
        with self.assertRaises(WrongPassword):
            BookFile.open(f.path, 'wrong')
        self.assertEqual(BookFile.open(f.path, 'secret').book['event']['name'], '示例赛事（虚构数据）')

    def test_versions_restore_and_log(self):
        f = BookFile.create(self.path, 'pw', tiny_book())
        f.save_version('1128版')
        f.replace_book(apply(f.book, [{'op': 'set', 'path': ['prices', 'x', 'A'], 'value': '999'}]), '修改')
        f.restore(f.versions()[0]['id'])
        g = BookFile.open(f.path, 'pw')
        self.assertEqual(g.book['prices']['x']['A'], '100')
        self.assertEqual([v['label'] for v in g.versions()], ['1128版', '恢复前自动保存'])
        self.assertEqual([e['action'] for e in g.data['log']][-1], '恢复版本')

    def test_change_password(self):
        f = BookFile.create(self.path, 'old', tiny_book())
        f.change_password('old', 'new')
        BookFile.open(f.path, 'new')
        with self.assertRaises(WrongPassword):
            BookFile.open(f.path, 'old')


class ExportTest(unittest.TestCase):
    def test_inventory_count_totals(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            out = inventory_count(tiny_book(), Path(d) / 'count.xlsx')
            rows = {r[0].strip(): r[1:] for r in load_workbook(out)['库存盘点'].iter_rows(values_only=True) if r and isinstance(r[0], str)}
            self.assertEqual(rows['销售门票总数量'], (450, 450))  # 300 sold + 150 priority
            self.assertEqual(rows['权益/赠票总数量'], (100, 100))
            self.assertEqual(rows['功能占用总数量'], (50, 50))
            self.assertEqual(rows['剩余未售'], (200, 200))
            self.assertEqual(rows['总座席'], (1000, 1000))
            forecast_table(example_book(), Path(d) / 'fc.xlsx')


class LiveTest(unittest.TestCase):
    def book(self):
        b = tiny_book()
        b['sessions'].append({'code': 'S2', 'date': '2030-01-01', 'band': 'x', 'layout': 'L'})
        b['sessions'].append({'code': 'S3', 'date': '2030-01-02', 'band': 'x', 'layout': 'L'})
        b['live'] = {
            'S1': {'at': '2030-01-01T15:00', 'checked': 900, 'realname': 950, 'total': 1000, 'female_pct': '90',
                   'local_pct': '20', 'age': {'18-24': '30'}, 'origins': [{'name': '甲省', 'pct': '50'}]},
            'S2': {'checked': 300, 'total': 500, 'female_pct': '60', 'age': {'18-24': '10'}},
        }
        return normalize(b)

    def test_session_day_and_event(self):
        b = self.book()
        live = live_summary(b, compute(b))
        s1 = live['sessions'][0]
        self.assertEqual((s1['rate'], s1['no_show'], s1['issued']), ('90.00', 100, 750))  # 100 comp + 150 + 200 + 300 sold
        self.assertEqual(s1['diff'], 250)
        self.assertFalse(live['sessions'][2]['entered'])
        day = live['days'][0]
        self.assertEqual((day['checked'], day['total'], day['rate']), (1200, 1500, '80.00'))
        self.assertEqual(day['female_pct'], '82.5')  # (90×900 + 60×300) / 1200
        self.assertEqual(day['age']['18-24'], '25.0')
        self.assertEqual(day['local_pct'], '20.0')  # only S1 gave it
        self.assertEqual(live['overall']['sessions'], 2)

    def test_bad_live_input_refused(self):
        with self.assertRaises(BookError):
            apply(self.book(), [{'op': 'set', 'path': ['live', 'S1', 'female_pct'], 'value': '120'}])
        with self.assertRaises(BookError):
            apply(self.book(), [{'op': 'set', 'path': ['live', 'S1', 'checked'], 'value': '-1'}])

    def test_session_rename_and_remove_follow(self):
        b = apply(self.book(), [{'op': 'rename', 'list': 'sessions', 'key': 'S1', 'to': 'S9'}])
        self.assertIn('S9', b['live'])
        b = apply(b, [{'op': 'remove', 'list': 'sessions', 'key': 'S9'}])
        self.assertNotIn('S9', b['live'])
        self.assertEqual(copy_event(b, 'N', 2031, 365)['live'], {})


ROW_PER_LINE = """项目ID\t项目名称\t项目时间\t城市
\t100000001\t示例杯【赞助商】\t2030-11-29至2030-12-07\t示例市\t项目结束\t数量（张）\t1,000\t900\t0\t100\t90%\t查看报表
\t\t\t\t\t金额（元）\t50000\t45000\t0\t5000
\t100000002\t示例杯\t2030-11-29至2030-12-07\t示例市\t销售中\t数量（张）\t2000\t500\t20\t1500\t25%\t查看报表
\t\t\t\t\t金额（元）\t20000000\t5000000.5\t2000\t14999999.5
"""
CELL_PER_LINE = "100000003\n示例杯\n2030-11-29至\n2030-12-07\n数量（张）\n123,450\n90000\n0\n33450\n72.90%\n金额（元）\n77777700\n55555500\n0\n22222200\n"


def screen_boxes(**skip):
    """OCR boxes laid out like Damai's on-site screen, with invented figures."""
    rows = [
        (300, 10, 200, 20, '示例混合团体赛 S3'), (330, 35, 150, 12, '2030-12-02 15:30:05'),
        (260, 60, 40, 15, '到场率'), (250, 80, 60, 20, '90.00%'),
        (30, 150, 70, 20, '已验票数'), (30, 175, 70, 22, '1,800'),
        (180, 150, 60, 18, '已实名数'), (180, 175, 70, 22, '1900人'),
        (280, 150, 50, 18, '总票数'), (280, 175, 70, 22, '2,000'),
        (420, 300, 60, 15, '观众性别'), (430, 400, 30, 12, '88%'), (520, 400, 30, 12, '12%'),
        (600, 300, 60, 15, '年龄分布'), (950, 300, 60, 15, '来源分布'), (1080, 298, 50, 15, '示例市'), (1140, 298, 30, 15, '21%'),
        (1000, 330, 50, 12, '甲省'), (1150, 330, 30, 12, '40%'),
    ]
    ages = [('18岁以下', '5%'), ('18-24岁', '30%'), ('25-29岁', '20%'), ('30-34岁', '15%'), ('35-39岁', '10%'),
            ('40-44岁', '8%'), ('45-49岁', '5%'), ('50岁以上', '7%')]
    for i, (label, value) in enumerate(ages):
        y = 330 + i * 26
        if label not in skip.get('labels', ()):
            rows.append((640, y, 50, 12, label))
        if value and label not in skip.get('values', ()):
            rows.append((830, y, 30, 12, value))
    return [Box(*r) for r in rows]


class ScreenshotTest(unittest.TestCase):
    def test_reads_every_figure(self):
        r = parse_boxes(screen_boxes())
        self.assertEqual(r['session'], 'S3')
        e = r['entry']
        self.assertEqual((e['at'], e['checked'], e['realname'], e['total']), ('2030-12-02T15:30', 1800, 1900, 2000))
        self.assertEqual((e['female_pct'], e['local_pct'], e['local_name']), ('88', '21', '示例市'))
        self.assertEqual(e['age'], {'u18': '5', '18-24': '30', '25-29': '20', '30-34': '15', '35-39': '10', '40-44': '8', '45-49': '5', '50+': '7'})
        self.assertEqual((r['missing'], r['warnings']), ([], []))

    def test_misread_label_still_placed_by_row(self):
        e = parse_boxes(screen_boxes(labels=('30-34岁',)))['entry']
        self.assertEqual(e['age']['30-34'], '15')

    def test_unread_value_left_empty_and_listed(self):
        r = parse_boxes(screen_boxes(values=('40-44岁',)))
        self.assertNotIn('40-44', r['entry']['age'])
        self.assertIn('age:40-44', r['missing'])

    def test_real_ocr_on_a_drawn_screen(self):
        """End to end with the OCR model, on a picture drawn from the invented layout (skipped without OCR or a Chinese font)."""
        from ticketbook import ocr
        font_path = next((f for f in ('/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', '/System/Library/Fonts/PingFang.ttc') if Path(f).exists()), None)
        if not ocr.available() or not font_path:
            self.skipTest('OCR or Chinese font not installed')
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new('RGB', (1220, 560), (10, 20, 70))
        draw = ImageDraw.Draw(image)
        for b in screen_boxes():
            draw.text((b.x, b.y), b.text, fill=(230, 235, 255), font=ImageFont.truetype(font_path, int(b.h)))
        with tempfile.TemporaryDirectory() as d:
            path = str(Path(d) / 'screen.png')
            image.save(path)
            r = ocr.read_screenshot(path)
        e = r['entry']
        self.assertEqual(r['session'], 'S3')
        self.assertEqual((e['checked'], e['realname'], e['total']), (1800, 1900, 2000))
        self.assertEqual(e.get('female_pct'), '88')

    def test_rate_mismatch_warns(self):
        boxes = [b for b in screen_boxes() if b.text != '1,800'] + [Box(30, 175, 70, 22, '1,600')]
        self.assertTrue(parse_boxes(boxes)['warnings'])


class DamaiTest(unittest.TestCase):
    def test_reads_copied_table(self):
        rows = parse_damai_table(ROW_PER_LINE)
        self.assertEqual([r['id'] for r in rows], ['100000001', '100000002'])
        self.assertEqual(rows[0]['name'], '示例杯【赞助商】')
        self.assertEqual((rows[0]['plan_qty'], rows[0]['sold_qty'], rows[0]['left_qty']), (1000, 900, 100))
        self.assertEqual((rows[1]['sold_amount'], rows[1]['left_amount']), ('5000000.5', '14999999.5'))

    def test_reads_one_cell_per_line(self):
        (row,) = parse_damai_table(CELL_PER_LINE)  # long amounts must not be taken for project IDs
        self.assertEqual((row['id'], row['plan_qty'], row['plan_amount']), ('100000003', 123450, '77777700'))

    def test_nothing_found(self):
        with self.assertRaises(BookError):
            parse_damai_table('随便一段文字 123')

    def test_snapshots_in_time_order(self):
        b = tiny_book()
        b['damai'] = [{'id': 'd2', 'at': '2030-11-20T10:00', 'projects': parse_damai_table(ROW_PER_LINE)},
                      {'id': 'd1', 'at': '2030-11-10T10:00', 'projects': parse_damai_table(CELL_PER_LINE)}]
        out = damai_summary(normalize(b), '10000000')
        self.assertEqual([s['id'] for s in out['snapshots']], ['d1', 'd2'])
        total = out['latest']['total']
        self.assertEqual((total['plan_qty'], total['sold_qty'], total['rate']), (3000, 1400, '46.67'))  # latest snapshot only
        self.assertEqual(total['sold_amount'], '5045000.5')
        self.assertEqual(out['vs_forecast'], '50.45')


class DesktopProtocolTest(unittest.TestCase):
    def test_round_trip(self):
        with tempfile.TemporaryDirectory() as d:
            s = Session()
            call = lambda m, **p: s.handle(json.dumps({'id': m, 'method': m, 'params': p}))
            self.assertEqual(call('state')['error']['code'], 'NO_FILE')
            created = call('create_book', path=str(Path(d) / 'a'), password='pw', example=True)
            self.assertTrue(created['ok'])
            self.assertTrue(created['result']['ledger']['ok'])
            r = call('forecast_apply', changes=[{'path': ['prices', 'pre', 'VIP'], 'value': '550'}])
            self.assertEqual(r['error']['code'], 'CONFIRM')
            r = call('forecast_apply', changes=[{'path': ['prices', 'pre', 'VIP'], 'value': '550'}], confirm=True)
            self.assertEqual(r['result']['book']['prices']['pre']['VIP'], '550')
            self.assertEqual(r['result']['versions'][0]['label'], '测算写回前自动保存')
            bad = call('edit', ops=[{'op': 'set', 'path': ['layouts', 'four', 'seats', 'VIP'], 'value': 'x'}])
            self.assertEqual(bad['error']['code'], 'INPUT')
            call('close')
            self.assertEqual(call('open_book', path=created['result']['path'], password='no')['error']['code'], 'PASSWORD')
            self.assertTrue(call('open_book', path=created['result']['path'], password='pw')['ok'])
            self.assertEqual(s.handle('not json')['error']['code'], 'PROTOCOL')
            self.assertEqual(len(call('damai_read', text=ROW_PER_LINE)['result']['projects']), 2)
            self.assertEqual(call('damai_add', text=ROW_PER_LINE)['error']['code'], 'INPUT')  # needs a time
            r = call('damai_add', text=ROW_PER_LINE, at='2030-11-20T10:00')
            self.assertEqual(r['result']['damai']['latest']['total']['sold_qty'], 1400)
            r = call('edit', ops=[{'op': 'set', 'path': ['live', 'S1'], 'value': {'checked': 10, 'total': 20}}])
            self.assertEqual(r['result']['live']['overall']['rate'], '50.00')


if __name__ == '__main__':
    unittest.main()
