"""Excel exports: the inventory count table (库存盘点) and the forecast table."""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from .ledger import LINE_NAMES, LINES, as_of, compute
from .forecast import forecast
from .model import BUCKET_KIND_NAMES, money, price


def _sheet_rows_inventory(book: dict, ledger: dict) -> list[list]:
    """Rows of the per-session inventory count: sales by round, then every bucket by kind, each block with a subtotal."""
    sessions = ledger['sessions']
    head = ['类别'] + [s['code'] for s in sessions] + ['合计']
    dates = [''] + [s['date'] for s in sessions] + ['']
    rows = [dates, head]

    def line(label, values):
        rows.append([label] + values + [sum(values)])
        return values

    def per_session(fn):
        return [fn(s) for s in sessions]

    def round_sold(code):
        return per_session(lambda s: sum(t['by_round'].get(code, 0) for t in s['tiers'].values()))

    def bucket_qty(bucket_id):
        return per_session(lambda s: sum(t['by_bucket'].get(bucket_id, 0) for t in s['tiers'].values()))

    rows.append(['销售门票'])
    sold = [0] * len(sessions)
    for rnd in book['rounds']:
        vals = line('  ' + rnd.get('name', rnd['code']), round_sold(rnd['code']))
        sold = [a + b for a, b in zip(sold, vals)]
    priority = [b for b in book['buckets'] if b['kind'] == 'priority']
    for b in priority:
        vals = line('  ' + b['name'], bucket_qty(b['id']))
        sold = [a + v for a, v in zip(sold, vals)]
    line('销售门票总数量', sold)
    for kind in ('comp', 'hold', 'reserve'):
        buckets = [b for b in book['buckets'] if b['kind'] == kind]
        if not buckets:
            continue
        rows.append([BUCKET_KIND_NAMES[kind]])
        subtotal = [0] * len(sessions)
        for b in buckets:
            vals = line('  ' + b['name'], bucket_qty(b['id']))
            subtotal = [a + v for a, v in zip(subtotal, vals)]
        line(BUCKET_KIND_NAMES[kind] + '总数量', subtotal)
    if book['products']:
        line('通票/套票占用', per_session(lambda s: s['total']['product']))
    line('剩余未售', per_session(lambda s: s['total']['left']))
    line('总座席', per_session(lambda s: s['total']['seats']))
    return rows


def _detail_rows(ledger: dict) -> list[list]:
    rows = [['场次', '日期', '票档', '票价'] + [LINE_NAMES[x] for x in LINES]]
    for s in ledger['sessions']:
        for tier, r in s['tiers'].items():
            rows.append([s['code'], s['date'], tier, money(r['price']) if r['price'] is not None else None] + [r[x] for x in LINES])
    return rows


def _price_rows(book: dict) -> list[list]:
    tiers = [t['code'] for t in book['tiers']]
    rows = [['价格段'] + tiers]
    for band in book['bands']:
        rows.append([band['name']] + [price(book, band['code'], t) for t in tiers])
    return rows


def _write(path: Path, sheets: list[tuple[str, list[list]]], title: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append([title])
        ws['A1'].font = Font(bold=True, size=13)
        for row in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
        ws.column_dimensions['A'].width = 22
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def inventory_count(book: dict, path, rounds: list[str] | None = None) -> Path:
    """The inventory count workbook. `rounds` limits sales to the rounds counted so far (None = all)."""
    view = as_of(book, rounds)
    ledger = compute(view)
    stamp = datetime.now().strftime('%Y-%m-%d %H:%M')
    title = f'{book["event"]["name"]} 票务数据明细（导出 {stamp}）'
    return _write(path, [('库存盘点', _sheet_rows_inventory(view, ledger)), ('分档明细', _detail_rows(ledger)), ('票价', _price_rows(book))], title)


def forecast_table(book: dict, path) -> Path:
    fc = forecast(book)
    rows = [['场次', '日期', '价格段', '中国队', '上座率%', '满座票房', '预计票房']]
    for s in fc['sessions']:
        rows.append([s['code'], s['date'], s['band'], '是' if s['china'] else '', money(s['fill']), money(s['full']), money(s['expected'])])
    for p in fc['products']:
        rows.append([p['name'], '', '', '', money(book['forecast']['product_fill']), money(p['full']), money(p['expected'])])
    rows += [[], ['预计总票房', money(fc['gross'])], ['代理费', money(fc['agent_fee'])], ['净票房', money(fc['net'])]]
    return _write(path, [('票房测算', rows)], f'{book["event"]["name"]} 票房测算')
