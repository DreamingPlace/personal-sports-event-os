"""Filling the event setup from the user's own planning files (Word plans, Excel sheets).

Tables are scanned for shapes that planning files use: a price table (票档 × 比赛阶段), seats per tier (总座席),
a session list (S1, S2… under their dates), and the event's name and dates from the text. Everything found is shown
for checking; nothing is written until the user confirms, and numbers are never guessed.
"""
from __future__ import annotations

import re
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .model import BookError

TIER = re.compile(r'^(S?VIP|Tier\s*\d+|[A-H]|[A-H]\s*档|[A-H]区|.{0,6}档|.{0,4}(遮挡|遮)|普通票|看台票|包厢)$', re.I)
HEADER_WORDS = {'票档', '档位', '票品', '档次', '价位', '票档/赛事'}


def is_tier(text: str) -> bool:
    return bool(text) and text not in HEADER_WORDS and bool(TIER.match(text))


STAGE = re.compile(r'预赛|资格赛|小组赛|循环赛|淘汰赛|1/4|八强|四强|半决赛|决赛|铜牌|首轮|次轮|第.轮|开幕|闭幕')
SESSION = re.compile(r'^S\d{1,3}$')
DATE = re.compile(r'(20\d\d)[-/.年](\d{1,2})[-/.月](\d{1,2})')


def _clean(value) -> str:
    if value is None:
        return ''
    if hasattr(value, 'isoformat') and not isinstance(value, str):
        return value.isoformat()[:10]
    return re.sub(r'\s+', '', str(value))


def _number(text: str):
    t = text.replace(',', '').replace('元', '').replace('¥', '')
    try:
        n = Decimal(t)
    except InvalidOperation:
        return None
    return n if n >= 0 else None


def _num_str(n: Decimal) -> str:
    return str(int(n)) if n == n.to_integral_value() else str(n.normalize())


# --- reading files into plain grids ---------------------------------------------------------------------------------
def read_grids(path) -> tuple[list[tuple[str, list[list[str]]]], list[str]]:
    """([(source name, grid)], text lines). Merged cells repeat their value, as a person reads them."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in ('.xlsx', '.xlsm'):
        return _xlsx_grids(p), []
    if suffix == '.docx':
        return _docx_grids(p)
    raise BookError('请选择 Word（.docx）或 Excel（.xlsx）文件')


def _xlsx_grids(path):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(str(path), data_only=True)
    except Exception as exc:  # openpyxl raises several types for broken files
        raise BookError('不是有效的 Excel 文件') from exc
    out = []
    for ws in wb.worksheets:
        rows = min(ws.max_row, 500)
        cols = min(ws.max_column, 80)
        grid = [[_clean(ws.cell(row=r, column=c).value) for c in range(1, cols + 1)] for r in range(1, rows + 1)]
        for rng in ws.merged_cells.ranges:
            value = _clean(ws.cell(row=rng.min_row, column=rng.min_col).value)
            for r in range(rng.min_row, min(rng.max_row, rows) + 1):
                for c in range(rng.min_col, min(rng.max_col, cols) + 1):
                    grid[r - 1][c - 1] = value
        out.append((ws.title, grid))
    return out


def _docx_grids(path):
    import docx
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # python-docx raises several types for files that are not Word documents
        raise BookError('不是有效的 Word（.docx）文件') from exc
    grids = []
    for i, table in enumerate(document.tables, 1):
        grids.append((f'表格{i}', [[_clean(c.text) for c in row.cells] for row in table.rows]))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return grids, lines


# --- finding things in the grids --------------------------------------------------------------------------------
def find_prices(source: str, grid: list[list[str]]) -> list[dict]:
    """Price tables. Long form: rows of [阶段, 票档, numbers…] with a header row naming each number column.
    Wide form: a header row of stage names over rows of [票档, numbers…]. Each usable number column is an option."""
    found = []
    for r, row in enumerate(grid):
        stage_cols = [c for c, v in enumerate(row) if v and STAGE.search(v) and len(v) <= 12]
        if len(stage_cols) < 2 or r + 1 >= len(grid):
            continue
        # wide: tiers down the first non-empty column under this header
        prices, tiers = {}, []
        for below in grid[r + 1:]:
            label = next((v for v in below if v), '')
            if not is_tier(label):
                if tiers:
                    break
                continue
            tiers.append(label)
            for c in stage_cols:
                n = _number(below[c]) if c < len(below) else None
                if n is not None:
                    prices.setdefault(row[c], {})[label] = _num_str(n)
        if len(tiers) >= 2 and prices:
            found.append({'source': source, 'kind': 'wide', 'options': [{'label': '、'.join(prices), 'prices': prices}], 'tiers': tiers})
    # long: a tier column with a stage column beside it
    for tc in range(max((len(r) for r in grid), default=0)):
        rows = [r for r in grid if tc < len(r) and is_tier(r[tc])]
        if len(rows) < 4:
            continue
        sc = next((c for c in range(tc) if sum(1 for r in rows if STAGE.search(r[c] or '')) >= len(rows) * 0.8), None)
        if sc is None:
            continue
        first = grid.index(rows[0])
        header = grid[first - 1] if first else []
        options = []
        for c in range(tc + 1, max(len(r) for r in rows)):
            prices = {}
            for r in rows:
                n = _number(r[c]) if c < len(r) else None
                if n is not None:
                    prices.setdefault(r[sc], {})[r[tc]] = _num_str(n)
            if sum(len(v) for v in prices.values()) >= len(rows) * 0.8:
                options.append({'label': header[c] if c < len(header) and header[c] else f'第{c + 1}列', 'prices': prices})
        if options:
            tiers = list(dict.fromkeys(r[tc] for r in rows))
            found.append({'source': source, 'kind': 'long', 'options': options, 'tiers': tiers})
    return found


def find_seats(source: str, grid: list[list[str]]) -> list[dict]:
    """Seats per tier: a row labelled 总座席 under a header row of tier names."""
    found = []
    for r, row in enumerate(grid):
        label = next((v for v in row if v), '')
        if '总座席' not in label:
            continue
        for h in range(r - 1, max(-1, r - 4), -1):
            header = grid[h]
            cols = [c for c, v in enumerate(header) if is_tier(v)]
            if len(cols) >= 2:
                seats = {}
                for c in cols:
                    n = _number(row[c]) if c < len(row) else None
                    if n is not None:
                        seats[header[c]] = int(n)
                if seats:
                    found.append({'source': source, 'seats': seats})
                break
    return found


def find_sessions(source: str, grid: list[list[str]]) -> list[dict]:
    """Sessions: a row of S1, S2… with their dates in a row above (a date may span several sessions)."""
    found = []
    for r, row in enumerate(grid):
        cols = [c for c, v in enumerate(row) if SESSION.match(v or '')]
        if len(cols) < 2:
            continue
        sessions = []
        for c in cols:
            day = ''
            for above in range(r - 1, max(-1, r - 4), -1):
                m = DATE.search(grid[above][c] if c < len(grid[above]) else '')
                if m:
                    day = date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
                    break
            sessions.append({'code': row[c], 'date': day})
        if any(s['date'] for s in sessions):
            found.append({'source': source, 'sessions': list({s['code']: s for s in sessions}.values())})
    return found


def find_event(lines: list[str]) -> dict:
    out = {}
    for line in lines[:5]:
        m = re.match(r'^(20\d\d)年?(.+(赛|杯|大满贯))', line)
        if m:
            out['year'], out['name'] = int(m.group(1)), m.group(2)
            break
    text = '\n'.join(lines)
    m = re.search(r'(\d{1,2})月(\d{1,2})日至(\d{1,2})月(\d{1,2})日', text)
    if m:
        out['dates'] = f'{m.group(1)}月{m.group(2)}日至{m.group(3)}月{m.group(4)}日'
    return out


def scan(paths: list[str]) -> dict:
    """Everything recognisable in the given files, for the user to check."""
    if not paths:
        raise BookError('请选择文件')
    result = {'prices': [], 'seats': [], 'sessions': [], 'event': {}}
    for path in paths:
        name = Path(path).name
        grids, lines = read_grids(path)
        for source, grid in grids:
            label = f'{name} · {source}'
            result['prices'] += find_prices(label, grid)
            result['seats'] += find_seats(label, grid)
            result['sessions'] += find_sessions(label, grid)
        for key, value in find_event(lines).items():
            result['event'].setdefault(key, value)
    return result


# --- turning a chosen finding into edits ------------------------------------------------------------------------
def ops_for(book: dict, choice: dict) -> list[dict]:
    """Edit operations for what the user ticked: {'prices': {...stage: {tier: price}}, 'seats': {...}, 'layout': code,
    'sessions': [...], 'event': {...}}. Tiers and stages are matched by name; new ones are added."""
    ops: list[dict] = []
    tiers = {t['name']: t['code'] for t in book['tiers']}
    bands = {b['name']: b['code'] for b in book['bands']}
    codes = {t['code'] for t in book['tiers']} | {b['code'] for b in book['bands']}

    def new_code(prefix, name):
        base = re.sub(r'[^0-9A-Za-z一-鿿]', '', name) or prefix
        code, i = base, 2
        while code in codes:
            code, i = f'{base}{i}', i + 1
        codes.add(code)
        return code

    def tier(name):
        if name not in tiers:
            tiers[name] = new_code('T', name)
            ops.append({'op': 'add', 'list': 'tiers', 'item': {'code': tiers[name], 'name': name}})
        return tiers[name]

    def band(name):
        if name not in bands:
            bands[name] = new_code('band', name)
            ops.append({'op': 'add', 'list': 'bands', 'item': {'code': bands[name], 'name': name}})
        return bands[name]

    for stage, row in (choice.get('prices') or {}).items():
        b = band(stage)
        for name, value in row.items():
            ops.append({'op': 'set', 'path': ['prices', b, tier(name)], 'value': value})
    seats = choice.get('seats') or {}
    layout = choice.get('layout') or ''
    if seats:
        if not any(lay['code'] == layout for lay in book['layouts']):
            layout = new_code('L', '导入布局')
            ops.append({'op': 'add', 'list': 'layouts', 'item': {'code': layout, 'name': '导入布局', 'seats': {}}})
        for name, n in seats.items():
            ops.append({'op': 'set', 'path': ['layouts', layout, 'seats', tier(name)], 'value': n})
    existing = {s['code'] for s in book['sessions']}
    default_layout = layout or (book['layouts'][0]['code'] if book['layouts'] else '')
    default_band = next(iter(bands.values()), '')
    for s in choice.get('sessions') or []:
        if s['code'] in existing:
            ops.append({'op': 'set', 'path': ['sessions', s['code'], 'date'], 'value': s['date']})
        else:
            ops.append({'op': 'add', 'list': 'sessions', 'item': {'code': s['code'], 'date': s['date'], 'start': '', 'band': default_band,
                                                                  'layout': default_layout, 'china': False, 'note': ''}})
    event = choice.get('event') or {}
    for key in ('name', 'year'):
        if event.get(key):
            ops.append({'op': 'set', 'path': ['event', key], 'value': event[key]})
    return ops
