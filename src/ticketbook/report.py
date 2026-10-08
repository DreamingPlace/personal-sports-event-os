"""Filling the user's own Word report templates with numbers from the ticket book.

A template is any .docx with markers such as {{每场可售}} or {{票价:预赛:VIP}} where a number should go.
A paragraph that holds only {{票价表}}, {{座席表}} or {{场次表}} is replaced by a table. The numbers always
come from the ticket book, never from AI. Markers the app does not know are left in place and listed,
so nothing is silently dropped. Templates are kept inside the encrypted ticket book file.
"""
from __future__ import annotations

import base64
import io
import re
from datetime import date
from decimal import ROUND_HALF_UP, Decimal

from .forecast import forecast
from .ledger import compute
from .model import BookError, index, money, price, product_prices

MARKER = re.compile(r'\{\{\s*([^{}]+?)\s*\}\}')
TABLES = ('票价表', '座席表', '场次表')
MAX_TEMPLATE = 20 * 1024 * 1024


def _cn_date(day: str) -> str:
    try:
        d = date.fromisoformat(day[:10])
    except ValueError:
        return day
    return f'{d.month}月{d.day}日'


def _num(value) -> str:
    """Whole numbers without decimals; others to 2 places, trailing zeros dropped."""
    d = Decimal(str(value))
    if d == d.to_integral_value():
        return str(int(d))
    return str(d.quantize(Decimal('0.01'), ROUND_HALF_UP)).rstrip('0').rstrip('.')


def _per_session(values: list[int]) -> str:
    """One number when every session has the same figure; otherwise the range, so a template never shows a wrong single value."""
    if not values:
        return '0'
    low, high = min(values), max(values)
    return str(low) if low == high else f'{low}–{high}'


def fields(book: dict) -> dict[str, str]:
    """Every marker name the app can fill, with its current value."""
    ledger, fc = compute(book), forecast(book)
    sessions = ledger['sessions']
    out: dict[str, str] = {
        '赛事名称': book['event']['name'] or '',
        '年份': str(book['event']['year'] or ''),
        '场馆': book['event']['venue'] or '',
        '场次数': str(len(sessions)),
    }
    days = sorted({s['date'] for s in book['sessions'] if s['date']})
    out['比赛天数'] = str(len(days))
    out['开始日期'] = _cn_date(days[0]) if days else ''
    out['结束日期'] = _cn_date(days[-1]) if days else ''

    def each(key):
        return [s['total'][key] for s in sessions]

    sale = [s['total']['sellable'] - s['total']['comp'] for s in sessions]  # 可售 as in the plans: after holds and comps
    for name, values in (('总座席', each('seats')), ('功能占用', each('hold')), ('权益票', each('comp')), ('优先购', each('priority')),
                         ('预留', each('reserve')), ('通票占用', each('product')), ('可售', sale), ('公开销售', each('public')),
                         ('已售', each('sold')), ('剩余', each('left'))):
        out['每场' + name] = _per_session(values)
        out['全程' + name] = str(sum(values))
    out['满场票房'] = _num(money(fc['full_value']))
    out['预计票房'] = _num(money(fc['gross']))
    out['预计票房万元'] = _num((money(fc['gross']) / 10000).quantize(Decimal('1'), ROUND_HALF_UP))
    out['代理费'] = _num(money(fc['agent_fee']))
    out['净票房'] = _num(money(fc['net']))
    bands, tiers = book['bands'], book['tiers']
    for band in bands:
        for tier in tiers:
            p = price(book, band['code'], tier['code'])
            out[f'票价:{band["name"]}:{tier["name"]}'] = '' if p is None else _num(p)
    for tier in tiers:
        out['座席:' + tier['name']] = _per_session([s['tiers'][tier['code']]['seats'] for s in sessions])
        out['可售:' + tier['name']] = _per_session([s['tiers'][tier['code']]['sellable'] - s['tiers'][tier['code']]['comp'] for s in sessions])
    for bucket in book['buckets']:
        out['分配:' + bucket['name']] = _per_session([sum(t['by_bucket'].get(bucket['id'], 0) for t in s['tiers'].values()) for s in sessions])
        caps = {t['code']: bucket.get('cap', {}).get(t['code']) for t in tiers}
        if any(v not in (None, '') for v in caps.values()):
            out['上限:' + bucket['name']] = str(sum(int(v) for v in caps.values() if v not in (None, '')))
            for tier in tiers:
                if caps[tier['code']] not in (None, ''):
                    out[f'上限:{bucket["name"]}:{tier["name"]}'] = str(caps[tier['code']])
    for p in book['products']:
        lines = product_prices(book, p)
        out[f'产品:{p["name"]}:数量'] = str(p['quota'])
        if lines:
            out[f'产品:{p["name"]}:票价'] = _per_session([int(unit) if unit == int(unit) else unit for _, _, unit in lines])
    session_names = {s['code']: s['code'] for s in book['sessions']}
    for rnd in book['rounds']:
        out[f'轮次:{rnd["name"]}:开售日期'] = _cn_date(rnd.get('opens') or '')
        out[f'轮次:{rnd["name"]}:场次'] = '全部场次' if rnd.get('sessions') is None else '、'.join(session_names.get(c, c) for c in rnd['sessions'])
        for tier in tiers:
            share = rnd['share'].get(tier['code'])
            if share not in (None, ''):
                out[f'轮次:{rnd["name"]}:{tier["name"]}:比例'] = _num(money(share)) + '%'
    return out


def _tables(book: dict) -> dict[str, list[list[str]]]:
    ledger, fc = compute(book), forecast(book)
    tiers = book['tiers']
    prices = [['比赛阶段'] + [t['name'] for t in tiers]]
    for band in book['bands']:
        row = [band['name']]
        for t in tiers:
            p = price(book, band['code'], t['code'])
            row.append('' if p is None else _num(p))
        prices.append(row)
    first = ledger['sessions'][0] if ledger['sessions'] else None
    seats = [['', *[t['name'] for t in tiers], '合计']]
    if first:
        def line(label, values):
            seats.append([label, *map(str, values), str(sum(values))])
        line('总座席', [first['tiers'][t['code']]['seats'] for t in tiers])
        for bucket in book['buckets']:
            values = [first['tiers'][t['code']]['by_bucket'].get(bucket['id'], 0) for t in tiers]
            if any(values):
                line(bucket['name'], values)
        line('可售', [first['tiers'][t['code']]['sellable'] - first['tiers'][t['code']]['comp'] for t in tiers])
    by_code = index(book['bands'])
    expected = {s['code']: s['expected'] for s in fc['sessions']}
    sessions = [['场次', '日期', '比赛阶段', '公开销售', '预计票房']]
    for s in ledger['sessions']:
        sessions.append([s['code'], s['date'], by_code.get(s['band'], {}).get('name', s['band']), str(s['total']['public']), _num(money(expected[s['code']]))])
    return {'票价表': prices, '座席表': seats, '场次表': sessions}


def load_template(raw: bytes):
    if len(raw) > MAX_TEMPLATE:
        raise BookError('模板文件过大（上限 20MB）')
    try:
        import docx
        return docx.Document(io.BytesIO(raw))
    except Exception as exc:  # python-docx raises several types for files that are not Word documents
        raise BookError('不是有效的 Word（.docx）文件') from exc


def _paragraphs(document):
    def walk(container):
        for p in container.paragraphs:
            yield p
        for table in getattr(container, 'tables', []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(document)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from walk(part)


def markers(raw: bytes) -> list[str]:
    found = []
    for p in _paragraphs(load_template(raw)):
        for name in MARKER.findall(p.text):
            if name not in found:
                found.append(name)
    return found


def evaluate(expression: str, values: dict[str, str]) -> str | None:
    """Simple sums in a marker, e.g. {{每场可售 - 5345}} or {{比赛天数 * 400}}: field names and numbers joined by
    + - * with spaces around each sign. Returns None when it is not such a sum or a value is not a single number."""
    parts = re.split(r'\s+([+\-*])\s+', expression)
    if len(parts) < 3:
        return None
    numbers = []
    for i, part in enumerate(parts):
        if i % 2:
            numbers.append(part)
            continue
        raw = values.get(part, part)
        try:
            numbers.append(Decimal(raw.rstrip('%')))
        except ArithmeticError:
            return None
    i = 0
    terms = [numbers[0]]  # multiply first, then add and subtract
    while i + 2 < len(numbers):
        op, value = numbers[i + 1], numbers[i + 2]
        if op == '*':
            terms[-1] = terms[-1] * value
        else:
            terms += [op, value]
        i += 2
    total = terms[0]
    for j in range(1, len(terms), 2):
        total = total + terms[j + 1] if terms[j] == '+' else total - terms[j + 1]
    return _num(total)


def check(book: dict, raw: bytes) -> dict:
    names, known = markers(raw), fields(book)
    return {'markers': names, 'unknown': [n for n in names if n not in known and n not in TABLES and evaluate(n, known) is None]}


def fill(book: dict, raw: bytes) -> tuple[bytes, list[str]]:
    """The filled document, and the markers that were left in place because the app does not know them."""
    document = load_template(raw)
    values, tables = fields(book), _tables(book)
    unknown: list[str] = []
    for p in list(_paragraphs(document)):
        text = p.text
        if '{{' not in text:
            continue
        only = MARKER.fullmatch(text.strip())
        if only and only.group(1) in tables:
            _insert_table(document, p, tables[only.group(1)])
            continue

        def swap(m):
            name = m.group(1)
            if name in values:
                return values[name]
            worked = evaluate(name, values)
            if worked is not None:
                return worked
            if name not in unknown:
                unknown.append(name)
            return m.group(0)

        runs = p.runs
        joined = ''.join(r.text for r in runs)
        for m in reversed(list(MARKER.finditer(joined))):
            new = swap(m)
            if new != m.group(0):
                replace_span(runs, m.start(), m.end(), new)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue(), unknown


def replace_span(runs, start: int, end: int, new: str):
    """Replace characters start..end of a paragraph's text, which may be spread over several runs.

    The new text takes the formatting of the run where the old text began; every other run keeps its own,
    so bold headings and coloured words around a marker are not lost.
    """
    offset = 0
    placed = False
    for run in runs:
        text = run.text
        lo, hi = offset, offset + len(text)
        offset = hi
        if hi <= start or lo >= end:
            continue
        cut_from, cut_to = max(start, lo) - lo, min(end, hi) - lo
        if not placed:
            run.text = text[:cut_from] + new + text[cut_to:]
            placed = True
        else:
            run.text = text[:cut_from] + text[cut_to:]


def _insert_table(document, paragraph, rows: list[list[str]]):
    table = document.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    try:
        table.style = 'Table Grid'
    except (KeyError, ValueError):
        pass  # templates without the built-in grid style still get a plain table
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            table.cell(i, j).text = value
    paragraph._p.addnext(table._tbl)
    paragraph._p.getparent().remove(paragraph._p)


def encode(raw: bytes) -> str:
    return base64.b64encode(raw).decode('ascii')


def decode(data: str) -> bytes:
    return base64.b64decode(data.encode('ascii'))
